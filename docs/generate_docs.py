import re
import os
import base64
import urllib.request
import zlib
import time
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def render_mermaid_to_png(mermaid_code, output_png_path):
    """Renders Mermaid code into a PNG image using mermaid.ink or kroki.io fallback"""
    os.makedirs(os.path.dirname(output_png_path), exist_ok=True)
    clean_code = mermaid_code.strip()
    
    # Try mermaid.ink first
    try:
        encoded = base64.urlsafe_b64encode(clean_code.encode('utf-8')).decode('ascii')
        url = f"https://mermaid.ink/img/{encoded}?bgColor=FFFFFF"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
        with urllib.request.urlopen(req, timeout=15) as res:
            if res.status == 200:
                with open(output_png_path, 'wb') as f:
                    f.write(res.read())
                return True
    except Exception:
        pass

    # Try kroki.io fallback
    try:
        compressed = zlib.compress(clean_code.encode('utf-8'), 9)
        kroki_encoded = base64.urlsafe_b64encode(compressed).decode('ascii')
        url = f"https://kroki.io/mermaid/png/{kroki_encoded}"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
        with urllib.request.urlopen(req, timeout=15) as res:
            if res.status == 200:
                with open(output_png_path, 'wb') as f:
                    f.write(res.read())
                return True
    except Exception:
        pass

    return False

def set_cell_background(cell, fill_color):
    """Sets background color of a table cell (HEX e.g. '1A365D')"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for table cells in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    """Applies clean subtle borders to a table"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_inline_runs(paragraph, text, base_font_size=10.5, is_italic_block=False):
    """Parses bold (**text**), code (`text`), italic (*text*), math ($math$), and links"""
    pattern = r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\$[^\$]+\$|\[[^\]]+\]\([^)]+\)|[^*`\$\[]+)'
    tokens = re.findall(pattern, text)
    
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
            run.bold = True
        elif token.startswith('$') and token.endswith('$'):
            math_text = token[1:-1].replace('\\text{', '').replace('}', '').replace('\\times', '×').replace('\\sum', 'Σ')
            run = paragraph.add_run(math_text)
            run.font.name = 'Cambria Math'
            run.italic = True
            run.bold = True
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
        elif token.startswith('[') and '](' in token:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                label = m.group(1)
                run = paragraph.add_run(label)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
                run.underline = True
            else:
                run = paragraph.add_run(token)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if is_italic_block:
                run.italic = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def convert_markdown_to_docx(md_path, docx_path, doc_title, doc_subtitle):
    doc = Document()
    
    # Page Setup (A4, 0.8 inch margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    # Header & Footer
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run(f"Enterprise E-Commerce SRS • Jahidul Islam (jahidcse181@gmail.com)")
    f_run.font.name = 'Calibri'
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(doc_title)
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D) # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run(doc_subtitle)
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68) # Slate Gray
    
    # Divider bar
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(16)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="2B6CB0"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(pBdr)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    diagrams_dir = os.path.join(os.path.dirname(md_path), "diagrams")
    os.makedirs(diagrams_dir, exist_ok=True)
    doc_base_name = os.path.splitext(os.path.basename(md_path))[0]
    mermaid_diagram_counter = 0

    in_code_block = False
    code_block_lang = ""
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_lines, code_block_lang, mermaid_diagram_counter
        if not code_lines:
            return
        code_text = "".join(code_lines).rstrip()
        
        # Check if this is a Mermaid diagram block -> Embed as Image!
        if code_block_lang.strip().lower() == 'mermaid':
            mermaid_diagram_counter += 1
            img_name = f"{doc_base_name}_diagram_{mermaid_diagram_counter}.png"
            img_path = os.path.join(diagrams_dir, img_name)
            
            # Ensure image exists
            if not os.path.exists(img_path) or os.path.getsize(img_path) == 0:
                render_mermaid_to_png(code_text, img_path)
            
            if os.path.exists(img_path) and os.path.getsize(img_path) > 0:
                # Add Visual Diagram Container
                diag_p = doc.add_paragraph()
                diag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                diag_p.paragraph_format.space_before = Pt(8)
                diag_p.paragraph_format.space_after = Pt(2)
                
                try:
                    with Image.open(img_path) as im:
                        w, h = im.size
                        max_w = 6.4
                        target_w = Inches(min(max_w, max(3.5, w / 150)))
                        doc.add_picture(img_path, width=target_w)
                except Exception:
                    doc.add_picture(img_path, width=Inches(6.2))
                
                # Center the inserted picture
                last_p = doc.paragraphs[-1]
                last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                last_p.paragraph_format.space_after = Pt(2)
                
                # Add Figure Caption
                caption_p = doc.add_paragraph()
                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_p.paragraph_format.space_after = Pt(12)
                c_run = caption_p.add_run(f"▲ Visual Flow Diagram #{mermaid_diagram_counter}")
                c_run.font.name = 'Calibri'
                c_run.font.size = Pt(8.5)
                c_run.italic = True
                c_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                
                code_lines = []
                code_block_lang = ""
                return

        # For normal code blocks (PHP, SQL, JSON, etc.) -> Render styled code container
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_table.autofit = False
        cell = box_table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F7FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Accent Border
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:color="CBD5E0"/><w:left w:val="single" w:sz="18" w:color="3182CE"/><w:bottom w:val="single" w:sz="6" w:color="CBD5E0"/><w:right w:val="single" w:sz="6" w:color="CBD5E0"/></w:tcBorders>')
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        lang_title = "CODE (" + code_block_lang.upper() + ")" if code_block_lang else "SPECIFICATION"
        badge_run = p.add_run(f"[{lang_title}]\n")
        badge_run.font.name = 'Consolas'
        badge_run.font.size = Pt(8)
        badge_run.bold = True
        badge_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        code_lines = []
        code_block_lang = ""

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        
        rows_data = []
        for line in table_lines:
            clean_line = line.strip()
            if not clean_line.startswith('|'):
                continue
            cells = [c.strip() for c in clean_line.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            rows_data.append(cells)
            
        if not rows_data:
            table_lines = []
            return

        num_cols = max(len(r) for r in rows_data)
        for r in rows_data:
            while len(r) < num_cols:
                r.append("")

        word_table = doc.add_table(rows=len(rows_data), cols=num_cols)
        word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(word_table, color="CBD5E0", sz="4")

        for r_idx, row_data in enumerate(rows_data):
            row = word_table.rows[r_idx]
            is_header = (r_idx == 0)
            
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if is_header:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for c_idx, cell_value in enumerate(row_data):
                cell = row.cells[c_idx]
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                
                if is_header:
                    set_cell_background(cell, "2B6CB0")
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F7FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if is_header:
                    run = p.add_run(cell_value)
                    run.bold = True
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    format_inline_runs(p, cell_value, base_font_size=9)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle Code / Diagram Blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            in_table = False
            flush_table()

        if stripped == '---':
            continue

        # Handle Math Display Blocks ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 2:
            math_content = stripped[2:-2].strip().replace('\\text{', '').replace('}', '').replace('\\times', ' × ').replace('\\sum', 'Σ ').replace('\\left(', '(').replace('\\right)', ')').replace('\\frac{', '').replace('}{', ' / ')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(f"✦  {math_content}")
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            continue

        # Handle Headings
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        elif stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(13.5)
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        elif stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        elif stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

        # Handle Bullet & Numbered Lists
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, stripped[2:].strip(), base_font_size=10)
        elif re.match(r'^\d+\.\s+', stripped):
            match = re.match(r'^(\d+\.)\s+(.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, match.group(2), base_font_size=10)

        # Handle Callouts / Blockquotes
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            format_inline_runs(p, stripped[2:].strip(), base_font_size=10, is_italic_block=True)

        # Standard Paragraph
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            format_inline_runs(p, stripped, base_font_size=10)

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    # Save with lock retry / fallback
    saved = False
    for attempt in range(3):
        try:
            doc.save(docx_path)
            print(f"[SUCCESS] Saved: {docx_path}")
            saved = True
            break
        except PermissionError:
            print(f"[LOCKED] File {docx_path} is currently open in Word. Retrying in 1s...")
            time.sleep(1)
            
    if not saved:
        # Fallback to an alternative file name if user has the original open in Word
        alt_path = docx_path.replace(".docx", "_UPDATED.docx")
        doc.save(alt_path)
        print(f"[SAVED AS ALTERNATIVE] Saved to: {alt_path} (Close original in Word to overwrite)")

if __name__ == "__main__":
    docs_dir = r"d:\xampp\htdocs\ORIO ECOMMERCE\docs"
    
    # 1. Developer SRS
    dev_md = os.path.join(docs_dir, "SRS_DEVELOPER_TECHNICAL_SPECIFICATION.md")
    dev_docx = os.path.join(docs_dir, "SRS_DEVELOPER_TECHNICAL_SPECIFICATION.docx")
    convert_markdown_to_docx(
        dev_md, 
        dev_docx, 
        "Software Requirements Specification (Developer Edition)", 
        "Technical Architecture, Database Schema, Anti-N+1 Strategy, Invoicing & Analytics Engine"
    )

    # 2. Client SRS
    client_md = os.path.join(docs_dir, "SRS_CLIENT_PROJECT_OVERVIEW.md")
    client_docx = os.path.join(docs_dir, "SRS_CLIENT_PROJECT_OVERVIEW.docx")
    convert_markdown_to_docx(
        client_md, 
        client_docx, 
        "Software Requirements Specification (Client Edition)", 
        "Executive Project Overview, Feature Catalog, Phased Roadmap & Business Flowcharts"
    )

    # 3. Development Tracker
    tracker_md = os.path.join(docs_dir, "DEVELOPMENT_STEP_BY_STEP_TRACKER.md")
    tracker_docx = os.path.join(docs_dir, "DEVELOPMENT_STEP_BY_STEP_TRACKER.docx")
    if os.path.exists(tracker_md):
        convert_markdown_to_docx(
            tracker_md, 
            tracker_docx, 
            "Development & Verification Tracker", 
            "Step-by-Step Implementation Roadmap, Module Breakdown & Testing Checkpoints"
        )

    # 4. Micro-Step Blueprint
    blueprint_md = os.path.join(docs_dir, "MICRO_STEP_DEVELOPMENT_BLUEPRINT.md")
    blueprint_docx = os.path.join(docs_dir, "MICRO_STEP_DEVELOPMENT_BLUEPRINT.docx")
    if os.path.exists(blueprint_md):
        convert_markdown_to_docx(
            blueprint_md, 
            blueprint_docx, 
            "Micro-Step Development & Testing Blueprint", 
            "Atomic 15-30 Minute Tasks with Immediate Execution & Verification Criteria"
        )

    # 5. 20-Phase Complete Roadmap
    roadmap_md = os.path.join(docs_dir, "20_PHASE_DEVELOPMENT_ROADMAP.md")
    roadmap_docx = os.path.join(docs_dir, "20_PHASE_DEVELOPMENT_ROADMAP.docx")
    if os.path.exists(roadmap_md):
        convert_markdown_to_docx(
            roadmap_md, 
            roadmap_docx, 
            "20-Phase Development & Verification Roadmap", 
            "Structured 20-Phase Implementation Blueprint with Milestones & Testing Checkpoints"
        )

    # 6. Master Blueprint (Conversation Guide)
    master_md = os.path.join(r"d:\xampp\htdocs\ORIO ECOMMERCE", "PROJECT_MASTER_BLUEPRINT.md")
    master_docx = os.path.join(docs_dir, "PROJECT_MASTER_BLUEPRINT.docx")
    if os.path.exists(master_md):
        convert_markdown_to_docx(
            master_md, 
            master_docx, 
            "Master Architecture & 20-Phase Conversation Blueprint", 
            "Permanent System Context, Architecture Rules, and Future Conversation Pick-up Guide"
        )






